# -*- coding: utf-8 -*-
"""P1 的金标准对照：成品形状（T17）与端到端可判分（T18）。

P1 是五个包里最小的（17 行）、也是唯一**六项门禁全过**的那一个，
所以它被选作 golden 层的基准包：

1. T17 —— 交出去的 xlsx 必须正好是 10 列 × 17 行、工作表名 ``响应偏离表``，
   表头与 ``config.COLUMNS`` 逐字相同，且**不能**多出标准答案里那一列
   ``判定说明``。这一条盯的是"成品与契约一致"，不是流水线内部状态。
2. T18 —— 同一个文件交给 ``grade.py`` 必须能解析，覆盖率 1.000，
   且混淆矩阵里**只有一格**判错。把那一格钉死，等于把"已知的那一个失误"
   写成可执行文档：以后任何人改动判定逻辑，只要这一格变了就会红。

与 agent.md 的偏差：T17 原文还要求"I1~I14 全通过"。这 14 条不变量已经在
``tests/invariants/test_output_invariants.py`` 里对**五个包统一**证明过，
在此重复只会是同一条能力的叠加，所以本文件只保留 P1 独有的部分
（行数、键集合、表头字面量），不变量交给 invariants 层。
"""

from __future__ import annotations

import pathlib

import pytest
from openpyxl import load_workbook

from bidtable import config
from bidtable.output import writer

pytestmark = [pytest.mark.golden]

PACK = "P1"
N_ROWS = 17


def _write_artifact(result, repo_root: pathlib.Path, dest: pathlib.Path) -> pathlib.Path:
    """按真实模板落盘，返回写出的路径。模板缺失时退回内置表头。"""
    template = repo_root / result.stats["template_xlsx"]
    written, warns = writer.write(
        result.rows, dest, template=template if template.is_file() else None
    )
    assert warns == [], warns
    return written


@pytest.mark.ec("J3")
@pytest.mark.case("T17")
def test_t17_the_delivered_p1_workbook_is_ten_columns_by_seventeen_rows(
    offline_run, repo_root, tmp_path
):
    """P1 的成品 xlsx：10 列、17 行、工作表 ``响应偏离表``、表头逐字对齐。

    边界在哪：列数是**闭集**，不是"至少 10 列"。标准答案有 11 列
    （多一列 ``判定说明``），所以最容易犯的错是为了对齐答案而多写一列——
    ``grade.py`` 用 ``.active`` 取表，多一列不会报错，但违反 I2。
    这里同时钉住"不能少"和"不能多"，以及行数正好 17（不是 18、不是 21）：
    21 是裁剪前的条目数，18 是折叠后的中间态，17 才是参与判定的行数。

    真实资料包实例：``训练集/P1``（钬激光治疗机 + 设备运维技术服务，2 个标的），
    用 ``03_响应偏离表_模板.xlsx`` 真实模板落盘。

    与 agent.md 的偏差：见模块 docstring —— I1~I14 不在此重复。
    """
    result = offline_run(PACK)
    written = _write_artifact(result, repo_root, tmp_path / f"{PACK}.xlsx")

    wb = load_workbook(written)
    try:
        # 只有一个工作表，且名字就是契约里的那个 —— grade.py 靠 .active 取表。
        assert wb.sheetnames == [config.SHEET_NAME]
        ws = wb.active
        assert ws.title == config.SHEET_NAME

        # 列是闭集：正好 10 列，逐字等于 config.COLUMNS，第 11 列必须不存在。
        assert ws.max_column == len(config.COLUMNS) == 10
        header = [ws.cell(1, i + 1).value for i in range(ws.max_column)]
        assert header == list(config.COLUMNS)
        assert "判定说明" not in header
        assert ws.cell(1, 11).value is None

        # 行数：1 行表头 + 17 行数据。
        assert ws.max_row == N_ROWS + 1
        assert len(result.rows) == N_ROWS
        # 裁剪链条的中间态不能泄漏到成品里。
        assert result.stats["cut_counts"]["抽取条目"] == 21
        assert result.stats["cut_counts"]["参与判定"] == N_ROWS
    finally:
        wb.close()

    # 成品与内存里的行必须一一对应：键集合相同，且键唯一（I1 在 P1 上的实例）。
    dev_col = list(config.COLUMNS).index("偏离类型") + 1
    lot_col = list(config.COLUMNS).index("标的") + 1
    sec_col = list(config.COLUMNS).index("分节") + 1
    no_col = list(config.COLUMNS).index("序号") + 1
    wb = load_workbook(written)
    try:
        ws = wb.active
        sheet_keys = [
            (str(ws.cell(r, lot_col).value), ws.cell(r, sec_col).value or "",
             str(ws.cell(r, no_col).value))
            for r in range(2, ws.max_row + 1)
        ]
        devs = {ws.cell(r, dev_col).value for r in range(2, ws.max_row + 1)}
    finally:
        wb.close()
    assert len(set(sheet_keys)) == len(sheet_keys) == N_ROWS
    assert sheet_keys == [(r.lot_no, r.section, r.no) for r in result.rows]
    assert devs <= set(config.DEVIATIONS)


@pytest.mark.ec("J1")
@pytest.mark.ec("J3")
@pytest.mark.case("T18")
def test_t18_p1_is_fully_covered_and_its_single_miss_is_pinned(score_pack, gold_map):
    """P1 端到端：覆盖率 1.000，``grade.py`` 可解析，唯一一格失误被钉死。

    边界在哪：这条测试同时守两端。
    一端是**格式**——``grade.py`` 是只读的评分器（O5），我们的文件必须能被它
    原样吃进去并吐出恰好 9 个键；键名或键数变了，说明成品结构漂了。
    另一端是**质量**——覆盖率必须是 1.000（17/17 一条不少），
    而偏离类型准确率 0.941 = 16/17，混淆矩阵里**有且只有一格**非对角：
    ``正偏离->无偏离``。把这一格写死，就等于把"我们知道错在哪"变成断言；
    将来若这一格消失（判定改好了）或变成别的（改坏了），测试都会红，
    逼着人回来更新这份记录，而不是让分数悄悄漂移。

    真实资料包实例：``训练集/P1`` 的 ``04_标准答案_响应偏离表.xlsx``。
    全程离线（无 Key），所以这条也顺带证明了 J1：降级通道不影响结构正确性。

    与 agent.md 的偏差：无。
    """
    result, score = score_pack(PACK)

    # grade.py 的输出契约：恰好这 9 个键，一个不多一个不少。
    assert set(score) == {
        "条目总数", "提交匹配条目", "多余条目", "覆盖率", "偏离类型准确率",
        "关键项风险识别", "证明材料命中率", "位置命中率", "混淆矩阵",
    }
    assert score["条目总数"] == N_ROWS
    assert score["提交匹配条目"] == N_ROWS
    assert score["多余条目"] == 0
    assert score["覆盖率"] == pytest.approx(1.0)
    assert score["偏离类型准确率"] == pytest.approx(16 / 17, abs=1e-3)
    assert score["证明材料命中率"] == pytest.approx(1.0)
    assert score["位置命中率"] == pytest.approx(1.0)
    # P1 只有一条 ★/▲ 风险行，且被正确识别 —— 这是它六项门禁全过的原因。
    assert score["关键项风险识别"].startswith("1/1")

    matrix = score["混淆矩阵"]
    diagonal = sum(v for k, v in matrix.items() if k.split("->")[0] == k.split("->")[1])
    assert diagonal == 16
    assert sum(matrix.values()) == N_ROWS
    off = {k: v for k, v in matrix.items() if k.split("->")[0] != k.split("->")[1]}
    assert off == {"正偏离->无偏离": 1}

    # 覆盖率 1.000 的结构性原因：我们的键集合与标准答案的键集合完全相同。
    gold_keys = set(gold_map(PACK))
    ours = {(r.lot_no, r.section, r.no) for r in result.rows}
    assert ours == gold_keys
