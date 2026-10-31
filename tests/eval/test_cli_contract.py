# -*- coding: utf-8 -*-
"""eval 层：交付命令行在**进程边界**上的契约（agent.md 6.8 的 C7 / 10.2 第 1 条）。

同一条命令在 unit 层已经跑过一遍（``tests/unit/test_run_cli.py`` 直接调 ``run.main``）。
进程内调用看不见四件事，只有真起子进程才看得见，所以本文件不复读 unit 层的断言，
只钉这四件：

1. **干净解释器**。pytest 会把仓库根塞进 ``sys.path``，于是"包没装好也能 import"
   这种问题在 unit 层永远不暴露。子进程只带 ``PYTHONPATH``，跟用户照 README 敲的一样。
2. **五个包全跑**。unit 层的双胞胎用例只用 P1；P3（71 行 / 二级分节）与
   P5（4 个标的 / C 族版面）的入口级回归只有这里挡得住。
3. **真实的管道**。重定向到 pipe 时 CPython 的 stdout 编码与缓冲都跟 tty 不同。
   Windows 上子进程的编码跟随**平台 ANSI 码本**（本机实测 cp936），不是 UTF-8；
   行尾也从 LF 变成 CRLF。今天退出码还是 0，仅仅因为 CLI 吐的每个字符
   （``｜ ★ ▲ ：``）恰好在 GBK 里都有。谁往进度或摘要里加一个码本外的字符
   （emoji、生僻字），一次**成功**的运行就会变成 ``UnicodeEncodeError`` traceback。
   第 2 条用例就是这道绊线，而且它显式指定窄码本，所以在 Linux CI 上同样有效。
4. **原地 vs 搬家**。unit 层比的是两份克隆（都不在原地），比不出"硬编码了
   ``训练集/P1``"；本文件比的是**原地真跑**与**搬到中文带空格目录、且 cwd 也换掉**
   的两次结果，两类路径假设一起抓。

★ 产物一律**按单元格**比，绝不按字节比：xlsx 是 zip，``docProps`` 里带时间戳，
  同一份输入两次跑出来的文件差 1 个字节（实测 8320 vs 8321）。

列宽、冻结窗格、底色属于 ``tests/unit/test_writer_schema.py``；
十列语义、取值域、位置格式属于 ``tests/invariants/``；
分数与六道门禁属于 ``tests/eval/test_selfscore.py``。本文件都不重复。
"""

from __future__ import annotations

import codecs
import shutil
import subprocess
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from bidtable import config

pytestmark = [pytest.mark.eval, pytest.mark.slow]

PACKS = ("P1", "P2", "P3", "P4", "P5")
# 摘要行报的是**输出行数**，P4 是 37（其中 1 行标准答案标了「不判分」，计分时算 36）。
OUTPUT_ROWS = {"P1": 17, "P2": 25, "P3": 71, "P4": 37, "P5": 56}
XLSX_DIMS = {"P1": "A1:J18", "P2": "A1:J26", "P3": "A1:J72", "P4": "A1:J38", "P5": "A1:J57"}
# 6.8 的 C7 要求目录名同时含中文字符与空格——两者走的是不同的转义路径。
_CJK_DIR = "中文 目录 名字"
_STAGES = tuple(f"[{i}/7]" for i in range(1, 8))
_EXCEL_ERRORS = ("#REF!", "#N/A", "#VALUE!", "#NAME?", "#DIV/0!", "#NULL!", "#NUM!")


def _grid(path: Path) -> tuple[str, str, list[list[Any]]]:
    """把工作簿摊成 (表名, 尺寸, 全部单元格值)。"""
    wb = load_workbook(path)
    ws = wb.active
    got = (ws.title, ws.dimensions, [[c.value for c in row] for row in ws.iter_rows()])
    wb.close()
    return got


@pytest.fixture(scope="module")
def runs(tmp_path_factory, repo_root, child_run, training_dir):
    """5 次原地真跑 + 1 次「搬到中文带空格目录、且换 cwd」的真跑。

    六次都带 ``--offline --no-cache``：离线让成绩不随机器上的 Key 状态飘，
    ``--no-cache`` 让门禁吃不到上一轮的旧产物（阶段缓存以文件名/大小/mtime 为键，
    代码改了而资料没改时它照样命中）。
    不带 ``--report``：复核报告的中文路径一致性 unit 层已经比过了。
    """
    base = tmp_path_factory.mktemp("cli")
    entry = str(repo_root / "run.py")          # 绝对路径：cwd 换了也得找得到入口
    inplace: dict[str, subprocess.CompletedProcess] = {}
    sheets: dict[str, Path] = {}
    for pack in PACKS:
        out = base / "inplace" / f"{pack}.xlsx"
        proc = child_run([entry, "--pack", str(training_dir / pack),
                          "--out", str(out), "--offline", "--no-cache"])
        if proc.returncode != 0:
            pytest.fail(f"{pack} 在干净解释器里退出码 {proc.returncode}\n{proc.stderr[-2000:]}")
        inplace[pack] = proc
        sheets[pack] = out

    cjk_pack = base / _CJK_DIR
    shutil.copytree(training_dir / "P1", cjk_pack)
    cjk = child_run([entry, "--pack", str(cjk_pack),
                     "--out", str(base / "cjk" / "P1.xlsx"),
                     "--offline", "--no-cache"], cwd=base)
    if cjk.returncode != 0:
        pytest.fail(f"搬到 {_CJK_DIR!r} 后退出码 {cjk.returncode}\n{cjk.stderr[-2000:]}")

    return {"inplace": inplace, "xlsx": sheets, "cjk": cjk, "entry": entry,
            "cjk_xlsx": base / "cjk" / "P1.xlsx", "training_dir": training_dir}


@pytest.mark.ec("J7")
@pytest.mark.ec("J8")
@pytest.mark.ec("J1")
def test_the_shipped_one_liner_exits_zero_for_every_pack_in_a_clean_interpreter(runs):
    """10.2 第 1 条：``run.py --pack <任意包> --out <任意路径>`` 五个包全部退出 0。

    边界在哪：**入口级**回归。前面几层都是从 ``bidtable.pipeline`` 往里调的，
    任何一处把 ``argparse`` 的默认值、``_validate`` 的取值域或者
    ``BidTableError`` → 退出码的映射改坏，那些层全绿而用户跑不起来。
    这里只查三件进程级的事：退出码、stdout 只有给人看的那一行摘要
    （进度一律走 stderr，所以 ``run.py … > summary.txt`` 是可用的）、
    以及落盘的 xlsx 能被重新打开成约定形状的工作簿。

    刻意**不**重抄 unit 层的内容：七个 ``[i/7]`` 阶段的文案、帮助分组、
    ``--version``、``_validate`` 的每条报错，都在 ``tests/unit/test_run_cli.py``。
    这里只断言"阶段行没有漏到 stdout"，不断言阶段行本身长什么样。
    """
    for pack in PACKS:
        proc = runs["inplace"][pack]
        assert "Traceback" not in proc.stderr, (pack, proc.stderr[-1500:])

        out_lines = proc.stdout.splitlines()
        assert len(out_lines) == 1, (pack, out_lines)
        assert out_lines[0].startswith(f"{pack}：{OUTPUT_ROWS[pack]} 行｜"), out_lines[0]
        assert not [s for s in _STAGES if s in proc.stdout], (pack, proc.stdout)

        # 10.2 第 4 条里机器可查的那部分：能重开、形状对、没有 Excel 错误哨兵。
        sheet = runs["xlsx"][pack]
        assert sheet.is_file(), sheet
        title, dims, grid = _grid(sheet)
        assert title == config.SHEET_NAME, (pack, title)
        assert dims == XLSX_DIMS[pack], (pack, dims)
        assert len(grid) == OUTPUT_ROWS[pack] + 1, pack          # 1 表头 + N 数据行
        assert {len(r) for r in grid} == {len(config.COLUMNS)}, pack
        leaked = [v for r in grid for v in r if isinstance(v, str) and v in _EXCEL_ERRORS]
        assert not leaked, (pack, leaked)


# Windows 简体中文 ANSI 码本。选它不是因为本机是中文 Windows，而是因为 CPython
# 在 Linux/macOS 上也内置这个码本 —— 显式指定之后，本用例在任何平台都不是空断言。
_NARROW = "cp936"


def _child_codec(child_run, extra_env: dict[str, str] | None = None) -> str:
    """让子进程**自己报出**它的 stdout 编码，而不是在测试里猜一个。"""
    probe = child_run(["-c", "import sys; print(sys.stdout.encoding)"],
                      utf8=False, extra_env=extra_env)
    assert probe.returncode == 0, probe.stderr.decode("utf-8", "replace")
    codec = probe.stdout.decode("ascii", "replace").strip()
    assert codec, "子进程没有报出 stdout 编码，无法验证码本兼容性"
    return codec


@pytest.mark.ec("J8")
def test_the_summary_line_survives_a_console_whose_codepage_is_not_utf8(runs, child_run, tmp_path):
    """控制台不是 UTF-8 时，一次**成功**的运行不能变成编码崩溃。

    边界在哪：子进程的 stdout 编码由**运行环境的码本**决定，不由源码里那行
    ``# -*- coding: utf-8 -*-`` 决定。同一行摘要，cp936 编出来 91 字节、
    UTF-8 编出来 122 字节，退出码都是 0——"都是 0"是运气，因为 CLI 目前吐的
    字符恰好全在 GBK 里。所以本用例**显式**把子进程钉到 cp936，再拿 cp936 去解码
    真跑的输出：解码不通过就等于用例失败，往输出里加码本外字符的人当场被拦下。
    逃生口是 ``PYTHONIOENCODING=utf-8``，已写进 README 已知问题。

    这里刻意不依赖"平台默认码本恰好是 cp936"：只报出它（``ambient``）用于失败信息，
    真正驱动断言的是显式指定的窄码本，于是 Linux CI 上跑到的也是同一条路径。
    """
    ambient = _child_codec(child_run)                        # 平台默认，仅作记录
    codec = _child_codec(child_run, {"PYTHONIOENCODING": _NARROW})
    # 码本名不能用字符串相等来比：CPython 会把它规整成 codecs.lookup(...).name，
    # 于是 PYTHONIOENCODING=cp936 的子进程报出来的是 'gbk'（同族的规范别名）。
    # 按规范名比，才不会被"请求的别名"与"报告的别名"这种差异假红。
    assert codecs.lookup(codec).name == codecs.lookup(_NARROW).name, (ambient, codec)

    proc = child_run([runs["entry"], "--pack", str(runs["training_dir"] / "P1"),
                      "--out", str(tmp_path / "cp.xlsx"), "--offline", "--no-cache"],
                     utf8=False, extra_env={"PYTHONIOENCODING": _NARROW})
    assert proc.returncode == 0, f"码本 {codec} 下退出码 {proc.returncode}：{proc.stderr[:400]!r}"
    assert b"Traceback" not in proc.stderr, proc.stderr[:800]

    text = proc.stdout.decode(codec)          # ← 码本容不下就会抛，正是本用例要抓的
    assert _STAGES[0].encode(codec) in proc.stderr, codec
    proc.stderr.decode(codec)                 # 进度行也得能在窄码本里活下来

    expected = runs["inplace"]["P1"].stdout
    # 非空守卫：窄码本的字节数必须跟 UTF-8 不同（中文在 cp936 是 2 字节、UTF-8 是 3 字节），
    # 否则说明码本压根没生效，本用例正在悄悄地什么也没测。
    assert len(proc.stdout) != len(expected.encode("utf-8")), (ambient, codec, len(proc.stdout))
    # 子进程在 Windows 上是文本模式，管道里是 CRLF；而 child_run(utf8=True) 走
    # universal-newlines，已经折成 LF。不归一化的话本用例在 Windows 上必红、
    # 在 Linux 上必绿 —— 那等于用平台差异冒充测试结果。
    assert text.replace("\r\n", "\n") == expected, (ambient, codec, text, expected)


@pytest.mark.ec("J5")
@pytest.mark.ec("J7")
@pytest.mark.case("B22")
def test_a_pack_relocated_into_a_cjk_and_space_path_is_cell_identical(runs):
    """C7（agent.md 6.8）：把资料包搬到含中文与空格的目录，结果必须**逐格相同**。

    边界在哪：路径在 Windows 上要经过 ``Path`` → ``str`` → ``fitz`` / ``openpyxl``
    → NTFS 四次转手，任何一处按本地码本编码就会静默写错文件或读不到文件；
    空格则会把"用 ``shell=True`` 拼命令行"或"按空格 split 参数"的实现打穿。
    本用例连 cwd 一起换到临时目录，于是"硬编码 ``训练集/P1``"与
    "依赖当前目录的相对路径"两类错误一起暴露。

    与 unit 层那条同名用例的分工：unit 比的是**两份克隆**（``plain`` vs 中文名），
    两边都不在原地，所以就算代码写死了 ``训练集/P1`` 它也照样绿；
    这里比的是**原地真跑 vs 搬家真跑**，才是真能抓住硬编码的那一版。
    报告文件不在这里比（unit 层已比过标题行与时间戳的差异）。

    ★ 只比单元格，不比字节：xlsx 内含时间戳，同一输入两次跑差 1 个字节。
    """
    summary = runs["cjk"].stdout.splitlines()
    assert len(summary) == 1, summary
    # 包名取自**目录名**，不是写死的 P1 —— 这一条才是真正的反硬编码断言。
    assert summary[0].startswith(f"{_CJK_DIR}：{OUTPUT_ROWS['P1']} 行｜"), summary[0]
    assert not summary[0].startswith("P1："), "包名被硬编码成 P1，而不是取目录名"
    assert "Traceback" not in runs["cjk"].stderr, runs["cjk"].stderr[-1500:]

    assert runs["cjk_xlsx"].is_file(), runs["cjk_xlsx"]
    assert _grid(runs["cjk_xlsx"]) == _grid(runs["xlsx"]["P1"])
