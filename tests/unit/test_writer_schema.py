"""``output/writer.py``：xlsx 列契约与写盘失败路径（I1~I3 / X24、X28、N4、N11、N12）。

七条用例各盯一个互不重叠的能力，不重复覆盖：

1. 列契约（X28）：严格 10 列、顺序 == ``COLUMNS``、表名 == ``响应偏离表``，
   并且**没有** gold 里的第 11 列 ``判定说明``。
2. 重复键自检（X24）：``dedupe`` 保留**第一条**、键做 ``strip``、空分节渲染成「分节空」。
3. 模板缺失（N4）：告警原文 + 用内置常量重建（列宽 / 加粗 / 底色 / freeze pane）。
4. 模板被改过：``_from_template`` 重命名活动表、删残留行、**静默**重写表头（一条告警都不发）。
5. 写盘失败（N12）：``PermissionError`` → 退出码 2 且消息点名「正被 Excel 打开」；
   普通 ``OSError`` 走另一条消息，两者可区分。
6. 输出目录不存在（N11）：``write`` 自己建父目录，中文与空格路径都要能落盘。
7. 换行开关 + 模板继承：``wrap=False`` 要真的落到单元格对齐；
   模板路径**不设** freeze pane（与内置重建路径不一致，这是实测行为，如实钉住）。
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from bidtable import errors
from bidtable.config import COLUMNS, SHEET_NAME
from bidtable.output import writer

pytestmark = [pytest.mark.unit]


@pytest.mark.ec("J3")
@pytest.mark.ec("J8")
@pytest.mark.case("X28")
def test_schema_is_exactly_ten_columns_in_template_order(mk_row, tmp_path):
    """X28：gold 有第 11 列「判定说明」，我们**绝不能**写出来。

    模板只有 ``A1:J1``，多写一列会让人工复核时列错位；
    ``grade.py`` 按表头名取列，所以顺序也必须与 ``COLUMNS`` 逐字一致。
    """
    row = mk_row(
        lot_no="1", section="一、主机", no="1.2", mark="★",
        req_text="输出功率：≥100W", resp_text="输出功率：120W",
        deviation="正偏离", evidence_file="YF_YF-9834Pro_技术规格书.docx",
        location="技术规格书 表行3",
    )
    out = tmp_path / "P1.xlsx"
    path, warnings = writer.write([row], out)
    assert path == out
    assert warnings == []

    wb = load_workbook(path)
    ws = wb.active
    assert ws.title == SHEET_NAME == "响应偏离表"
    # 先读尺寸再读单元格：openpyxl 的 ws.cell() 会**创建**被访问的格子，
    # 先访问第 11 列再把 max_column 断言成 10 是自己骗自己。
    assert ws.max_column == len(COLUMNS) == 10
    assert ws.max_row == writer.FIRST_DATA_ROW == 2   # 数据从第 2 行开始
    assert ws.dimensions == "A1:J2"             # ← 物理上就没有第 11 列
    header = [ws.cell(row=1, column=i).value for i in range(1, 11)]
    assert header == list(COLUMNS)
    assert "判定说明" not in header
    assert [ws.cell(row=2, column=i).value for i in range(1, 11)] == row.cells()
    wb.close()


@pytest.mark.ec("J8")
@pytest.mark.ec("J3")
@pytest.mark.case("X24")
def test_dedupe_keeps_the_first_row_and_names_the_key(mk_row):
    """X24 / N14：``grade.py`` 把行装进 dict，重复键会**静默覆盖**，丢的是整行所有指标。

    所以写出前必须自检。三个边界点一起钉住：
    键是 ``strip`` 过的（``" 1 "`` 与 ``"1"`` 同键）、保留的是**先出现**的那条、
    空分节在告警里渲染成「分节空」（否则会打印出 ``分节,`` 这种读不懂的片段）。
    """
    first = mk_row(lot_no=" 1 ", section="", no="2", req_text="第一条")
    second = mk_row(lot_no="1", section="  ", no=" 2 ", req_text="第二条")
    other = mk_row(lot_no="1", section="", no="3", req_text="第三条")

    kept, warnings = writer.dedupe([first, second, other])
    assert kept == [first, other]
    assert first.req_text == "第一条"            # 保留先出现的
    assert warnings == [
        "输出现重复键 (标的1, 分节空, 序号2)，"
        "已丢弃后出现的那一行（grade.py 的行字典会静默覆盖重复键）。"
    ]
    assert writer.dedupe([first, other]) == ([first, other], [])


@pytest.mark.ec("J1")
@pytest.mark.ec("J7")
@pytest.mark.ec("J8")
@pytest.mark.case("N4")
def test_missing_template_warns_and_rebuilds_from_builtin_constants(mk_row, tmp_path):
    """N4：``scope.json`` 里的模板路径是**相对**路径，换个 cwd 跑就会找不到。

    这时不能崩，也不能静默：发一条指名道姓的告警，然后用实测出来的样式常量
    （列宽 / 加粗 / 底色 ``E8EEF4`` / ``freeze_panes="A2"``）重建，产物照样能判分。
    """
    tpl = tmp_path / "根本没有这个模板.xlsx"
    out = tmp_path / "P1.xlsx"
    path, warnings = writer.write([mk_row(no="1")], out, template=tpl)
    assert path == out and out.is_file()
    assert warnings == [f"模板 {tpl} 不存在，已按内置样式常量重建表头与列宽。"]

    wb = load_workbook(path)
    ws = wb.active
    for idx, width in enumerate(writer.WIDTHS, start=1):
        assert ws.column_dimensions[get_column_letter(idx)].width == width
    a1 = ws["A1"]
    assert a1.value == COLUMNS[0]
    assert a1.font.bold is True
    assert str(a1.fill.start_color.rgb).endswith(writer.HEADER_FILL)
    assert a1.fill.fill_type == "solid"
    assert ws.freeze_panes == f"A{writer.FIRST_DATA_ROW}"
    wb.close()


@pytest.mark.ec("J3")
@pytest.mark.ec("J8")
def test_tampered_template_is_repaired_silently(mk_row, tmp_path):
    """模板被改过（表名不对、表头不对、还残留示例行）时的自愈路径。

    ``_from_template`` 重命名活动表、删掉残留行、以 ``COLUMNS`` 为准重写表头，
    **而且一条告警都不发**。这个「静默」是实测行为，如实钉住：
    只有「模板文件不存在」那条路径会告警，表头对不上属于安静修复。
    """
    tpl = tmp_path / "被改过的模板.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for i in range(1, 11):
        ws.cell(row=1, column=i, value=f"错{i}")
    for r in range(2, 5):
        ws.cell(row=r, column=1, value=f"残留{r}")
    wb.save(tpl)
    wb.close()

    out = tmp_path / "P1.xlsx"
    _, warnings = writer.write([mk_row(no="1"), mk_row(no="2")], out, template=tpl)
    assert warnings == []

    wb = load_workbook(out)
    ws = wb.active
    assert ws.title == SHEET_NAME
    assert [ws.cell(row=1, column=i).value for i in range(1, 11)] == list(COLUMNS)
    assert ws.max_row == 3                      # 残留 3 行被删，只剩表头 + 2 行数据
    # 第 4 列是「序号」；第 1 列是「标的」，两行都是 "1"，
    # 残留行写在第 1 列，所以顺带断言第 1 列里再也没有「残留」字样。
    assert [ws.cell(row=2, column=4).value, ws.cell(row=3, column=4).value] == ["1", "2"]
    assert all("残留" not in str(ws.cell(row=r, column=1).value) for r in (2, 3))
    wb.close()


@pytest.mark.ec("J7")
@pytest.mark.ec("J5")
@pytest.mark.case("N12")
def test_write_failures_become_output_error_with_exit_code_two(mk_row, tmp_path, monkeypatch):
    """N12：目标被 Excel 占用是这套工具最常见的现场事故。

    ``PermissionError`` 要转成退出码 2 的 :class:`OutputError`，消息点名
    「正被 Excel 打开」并带上完整路径；普通 ``OSError``（磁盘满等）走另一条消息。
    两条都必须可区分、都不留半成品文件。
    """
    out = tmp_path / "P1.xlsx"
    rows = [mk_row(no="1")]

    def _locked(self, path):
        raise PermissionError(13, "Permission denied", str(path))

    monkeypatch.setattr(writer.Workbook, "save", _locked)
    with pytest.raises(errors.OutputError) as ei:
        writer.write(rows, out)
    assert isinstance(ei.value, errors.BidTableError)
    assert ei.value.exit_code == 2
    assert "正被 Excel 打开" in str(ei.value)
    assert str(out) in str(ei.value)
    assert not out.exists()

    def _full(self, path):
        raise OSError(28, "No space left on device")

    monkeypatch.setattr(writer.Workbook, "save", _full)
    with pytest.raises(errors.OutputError) as ei2:
        writer.write(rows, out)
    assert ei2.value.exit_code == 2
    assert "No space left on device" in str(ei2.value)
    assert "正被 Excel 打开" not in str(ei2.value)     # 两条分支不串味


@pytest.mark.ec("J5")
@pytest.mark.ec("J7")
@pytest.mark.case("N11")
def test_output_path_in_a_missing_nested_directory_is_created(mk_row, tmp_path):
    """N11：``--out`` 指到一个还不存在的多级目录时必须自己建父目录。

    顺手把 B22 的形态也压进来：路径里同时有中文和空格
    （``输出/P1 结果/响应偏离表.xlsx``），这是评审现场最真实的落盘位置。
    """
    out = tmp_path / "输出" / "P1 结果" / "响应偏离表.xlsx"
    assert not out.parent.exists()
    path, warnings = writer.write([mk_row(no="1")], out)
    assert warnings == []
    assert path == out and out.is_file()
    wb = load_workbook(path)
    assert wb.active.max_row == 2
    wb.close()


@pytest.mark.ec("J8")
@pytest.mark.ec("J3")
def test_wrap_flag_reaches_alignment_and_template_keeps_its_own_geometry(
    mk_row, tmp_path, pack_path
):
    """换行开关要真的落到单元格对齐上，且两条构造路径的几何属性**并不完全一致**。

    ``wrap=False`` 给窄终端导出用；模板路径继承模板自己的列宽，
    但 ``_from_template`` **不设** ``freeze_panes``，而内置重建路径会设 ``A2``。
    模块 docstring 说「两条路径长得一样」，实测差这一个属性，如实钉住并记进已知问题。
    """
    rows = [mk_row(no="1", req_text="很长的一行招标要求原文")]
    plain = tmp_path / "plain.xlsx"
    writer.write(rows, plain)
    nowrap = tmp_path / "nowrap.xlsx"
    writer.write(rows, nowrap, wrap=False)

    wb = load_workbook(plain)
    assert wb.active["F2"].alignment.wrap_text is True
    assert wb.active["F2"].alignment.vertical == "top"
    assert wb.active.freeze_panes == "A2"
    wb.close()

    wb = load_workbook(nowrap)
    # openpyxl 不序列化「等于默认值」的属性，wrap_text=False 存盘后再读回来是 None。
    # 所以要断言的是**假值**，而不是字面 False。
    assert not wb.active["F2"].alignment.wrap_text
    assert wb.active["F2"].alignment.vertical == "top"
    wb.close()

    tpl = pack_path("P1") / "03_响应偏离表_模板.xlsx"
    if not tpl.is_file():
        pytest.skip(f"缺模板 {tpl}")
    from_tpl = tmp_path / "from_tpl.xlsx"
    writer.write(rows, from_tpl, template=tpl)
    wb = load_workbook(from_tpl)
    ws = wb.active
    for idx, width in enumerate(writer.WIDTHS, start=1):
        assert ws.column_dimensions[get_column_letter(idx)].width == width
    assert ws.freeze_panes is None              # ← 模板路径不补 freeze pane
    assert ws["F2"].alignment.wrap_text is True
    wb.close()
